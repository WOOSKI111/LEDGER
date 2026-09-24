import os
import io
from flask import Flask, render_template, request, jsonify, session, send_file, redirect, url_for
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import db
import ocr

app = Flask(__name__)
app.secret_key = os.environ.get("SECRET_KEY", "dev-key-change-this-in-production")

db.init_db()


@app.route("/")
def index():
    balances = db.get_balances()
    total_owed = sum(b["owed"] or 0 for b in balances)
    recent = db.get_invoices()[:8]
    return render_template("index.html", balances=balances, total_owed=total_owed, recent=recent)


@app.route("/scan", methods=["POST"])
def scan():
    if "invoice" not in request.files:
        return jsonify({"error": "No image uploaded"}), 400

    file = request.files["invoice"]
    media_type = file.mimetype
    if media_type not in ("image/jpeg", "image/png", "image/webp"):
        media_type = "image/jpeg"

    image_bytes = file.read()
    if not image_bytes:
        return jsonify({"error": "Empty image"}), 400

    try:
        header, items = ocr.extract_invoice(image_bytes, media_type)
    except Exception as e:
        return jsonify({"error": f"Extraction failed: {str(e)}"}), 500

    if not items:
        return jsonify({"error": "Couldn't find any line items on that invoice. Try a clearer photo, or add it manually."}), 200

    session["pending_header"] = header
    session["pending_items"] = items
    return jsonify({"header": header, "items": items})


@app.route("/review")
def review():
    header = session.get("pending_header", {})
    items = session.get("pending_items", [])
    return render_template("review.html", header=header, items=items)


@app.route("/review/manual")
def review_manual():
    return render_template("review.html", header={}, items=[{"part_name": "", "part_number": "", "cost": "", "fits": ""}])


@app.route("/confirm", methods=["POST"])
def confirm():
    data = request.get_json()
    header = data.get("header", {})
    items = [i for i in data.get("items", []) if i.get("part_name", "").strip()]
    if not header.get("supplier", "").strip() or not items:
        return jsonify({"error": "Need a supplier and at least one part."}), 400
    invoice_id = db.insert_invoice(header, items)
    session.pop("pending_header", None)
    session.pop("pending_items", None)
    return jsonify({"ok": True, "invoice_id": invoice_id})


@app.route("/invoices")
def invoices():
    supplier = request.args.get("supplier", "").strip() or None
    search = request.args.get("search", "").strip() or None
    status = request.args.get("status", "").strip()
    paid = {"paid": True, "unpaid": False}.get(status)
    entries = db.get_invoices(supplier=supplier, paid=paid, search=search)
    return render_template("invoices.html", invoices=entries, supplier=supplier or "", search=search or "", status=status)


@app.route("/invoice/<int:invoice_id>/paid", methods=["POST"])
def toggle_paid(invoice_id):
    paid = request.get_json().get("paid", True)
    db.set_paid(invoice_id, paid)
    return jsonify({"ok": True})


@app.route("/invoice/<int:invoice_id>/delete", methods=["POST"])
def delete_invoice(invoice_id):
    db.delete_invoice(invoice_id)
    return redirect(url_for("invoices"))


@app.route("/invoice/<int:invoice_id>/update", methods=["POST"])
def update_invoice(invoice_id):
    db.update_invoice(invoice_id, request.get_json())
    return jsonify({"ok": True})


@app.route("/item/<int:item_id>/update", methods=["POST"])
def update_item(item_id):
    db.update_line_item(item_id, request.get_json())
    return jsonify({"ok": True})


@app.route("/item/<int:item_id>/delete", methods=["POST"])
def delete_item(item_id):
    db.delete_line_item(item_id)
    return jsonify({"ok": True})


@app.route("/export")
def export():
    entries = db.get_invoices()

    wb = Workbook()

    # Sheet 1: invoice summary
    ws1 = wb.active
    ws1.title = "Invoices"
    headers1 = ["Supplier", "Invoice #", "Invoice Date", "Due Date", "Total", "Status", "Paid Date", "Notes"]
    ws1.append(headers1)
    _style_header(ws1, headers1)
    for inv in entries:
        ws1.append([
            inv["supplier"], inv["invoice_number"], inv["invoice_date"], inv["due_date"],
            inv["total_cost"], "Paid" if inv["paid"] else "Unpaid", inv["paid_date"], inv["notes"],
        ])
    _autofit(ws1, headers1)

    # Sheet 2: full line-item detail
    ws2 = wb.create_sheet("Line Items")
    headers2 = ["Supplier", "Invoice #", "Part Name", "Part Number", "Cost", "Fits", "Status"]
    ws2.append(headers2)
    _style_header(ws2, headers2)
    for inv in entries:
        for item in inv["items"]:
            ws2.append([
                inv["supplier"], inv["invoice_number"], item["part_name"], item["part_number"],
                item["cost"], item["fits"], "Paid" if inv["paid"] else "Unpaid",
            ])
    _autofit(ws2, headers2)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name="ledger_export.xlsx",
                      mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def _style_header(ws, headers):
    fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill


def _autofit(ws, headers):
    for i, header in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(14, len(header) + 4)


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
